"""Import "The Ugly Tracker" workbook (both tabs). Idempotent: re-running updates, never duplicates."""
import sys
from datetime import date, datetime, timedelta

import openpyxl

from . import db
from .logic import monday

PALETTE = ["#ef4444", "#f59e0b", "#10b981", "#3b82f6", "#a855f7", "#ec4899", "#14b8a6", "#f97316"]


def _find_header(rows, first_cell):
    for i, r in enumerate(rows):
        if r and str(r[0] or "").strip().lower() == first_cell.lower():
            return i
    return None


def _num(v, default=0):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return default


def _clean(v):
    s = str(v).strip() if v is not None else ""
    return "" if s in ("—", "-", "–") else s


def _parse_date(v):
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    for fmt in ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(v).strip(), fmt).date()
        except ValueError:
            pass
    return None


def _split_vsc(total):
    """The workbook only stores the /30 total; split it evenly across V/S/C."""
    total = max(0, min(30, total))
    base, extra = divmod(total, 3)
    return [min(10, base + (1 if i < extra else 0)) for i in range(3)]


def import_workbook(ws: int, path, week_start: date | None = None) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False)
    week_start = week_start or monday(date.today())
    counts = {"phones": 0, "accounts": 0, "views": 0, "posts": 0}
    sheets = {ws.title.lower(): ws for ws in wb}

    tracker = next((ws for t, ws in sheets.items() if "view" in t), None)
    if tracker is not None:
        rows = list(tracker.iter_rows(values_only=True))
        h = _find_header(rows, "Phone")
        if h is not None:
            phone_ids = {}
            for r in rows[h + 1:]:
                if not r[0] or str(r[0]).strip().upper().startswith("TOTALS") or not r[1]:
                    break
                n = _num(r[0])
                if n not in phone_ids:
                    row = db.one("SELECT id FROM phones WHERE workspace_id = %s AND label = %s", (ws, f"Phone {n}"))
                    if row is None:
                        row = db.one(
                            "INSERT INTO phones(workspace_id, label, color, position) VALUES (%s,%s,%s,%s) RETURNING id",
                            (ws, f"Phone {n}", PALETTE[(n - 1) % len(PALETTE)], n))
                        counts["phones"] += 1
                    phone_ids[n] = row["id"]
                times = _num(r[5], 1)
                status = "paused" if times == 0 else "active"
                acc = db.one(
                    "INSERT INTO accounts(workspace_id, phone_id, handle, platform, niche, status, times_per_day, content_hook, model) "
                    "VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s) ON CONFLICT (workspace_id, handle) DO UPDATE SET "
                    "phone_id=EXCLUDED.phone_id, platform=EXCLUDED.platform, niche=EXCLUDED.niche, "
                    "times_per_day=EXCLUDED.times_per_day, content_hook=EXCLUDED.content_hook, model=EXCLUDED.model "
                    "RETURNING id",
                    (ws, phone_ids[n], str(r[1]).strip(), _clean(r[3]) or "TikTok", _clean(r[4]),
                     status, times, _clean(r[6]), _clean(r[7])))
                counts["accounts"] += 1
                # Day 1..7 sit in columns J, L, N, P, R, T, V (every other column from index 9)
                for i in range(7):
                    v = r[9 + 2 * i] if len(r) > 9 + 2 * i else None
                    if v is None:
                        continue
                    db.execute(
                        "INSERT INTO daily_views(account_id, workspace_id, day, views, source) VALUES (%s,%s,%s,%s,'manual') "
                        "ON CONFLICT (account_id, day) DO UPDATE SET views=EXCLUDED.views, updated_at=now()",
                        (acc["id"], ws, week_start + timedelta(days=i), _num(v)))
                    counts["views"] += 1

    hooks = next((ws for t, ws in sheets.items() if "hook" in t), None)
    if hooks is not None:
        rows = list(hooks.iter_rows(values_only=True))
        h = _find_header(rows, "S.No")
        if h is not None:
            for r in rows[h + 1:]:
                if r[0] is None or not r[2]:
                    break
                d = _parse_date(r[1])
                handle = str(r[2]).strip()
                if d is None:
                    continue
                acc = db.one("SELECT id FROM accounts WHERE workspace_id = %s AND handle = %s", (ws, handle))
                if acc is None:
                    acc = db.one("INSERT INTO accounts(workspace_id, handle) VALUES (%s,%s) RETURNING id", (ws, handle))
                    counts["accounts"] += 1
                hook = _clean(r[4])
                if db.one("SELECT 1 FROM posts WHERE workspace_id=%s AND account_id=%s AND post_date=%s AND hook=%s",
                          (ws, acc["id"], d, hook)):
                    continue
                vsc = _num(r[10], -1)
                v, s, c = _split_vsc(vsc) if vsc >= 0 else (None, None, None)
                db.execute(
                    "INSERT INTO posts(workspace_id, account_id, post_date, status, ai_prompt, hook, body, caption, views, "
                    "comments, intent, v, s, c, score_notes, score_by) "
                    "VALUES (%s,%s,%s,'posted',%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                    (ws, acc["id"], d, _clean(r[3]), hook, _clean(r[5]), _clean(r[6]), _num(r[7]), _num(r[8]),
                     str(r[9] or "").strip().lower() == "yes", v, s, c,
                     "Imported total VSC split evenly across V/S/C" if v is not None else "",
                     "imported" if v is not None else ""))
                counts["posts"] += 1
    return counts


if __name__ == "__main__":
    db.init_schema()
    if len(sys.argv) < 3:
        sys.exit("usage: python -m app.import_xlsx <workspace_id> <file.xlsx> [week_start YYYY-MM-DD]")
    print(import_workbook(int(sys.argv[1]), sys.argv[2],
                          date.fromisoformat(sys.argv[3]) if len(sys.argv) > 3 else None))
